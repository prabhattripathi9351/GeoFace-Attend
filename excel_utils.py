from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
import io

def generate_attendance_excel(matched_students, subject="General"):
    """Matched students se Excel file banata hai"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    # Title
    ws.merge_cells('A1:H1')
    title = ws['A1']
    title.value = "📊 Attendance Report"
    title.font = Font(bold=True, size=14, color="FFFFFF")
    title.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Info
    ws.merge_cells('A2:H2')
    info = ws['A2']
    now = datetime.now()
    info.value = f"Date: {now.strftime('%d %b %Y')} | Time: {now.strftime('%H:%M')} | Present: {len(matched_students)}"
    info.font = Font(italic=True, size=10)
    info.alignment = Alignment(horizontal="center")

    # Headers
    headers = ['S.No', 'Roll No', 'Name', 'Course', 'Branch', 'Year', 'Status', 'Time']
    ws.append(headers)

    fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    font = Font(bold=True, color="FFFFFF")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for cell in ws[3]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    # Data
    for i, m in enumerate(matched_students, 1):
        s = m['student']
        ws.append([
            i,
            s.get('roll_no', 'N/A'),
            s.get('name', 'N/A'),
            s.get('course', 'N/A'),
            s.get('branch', 'N/A'),
            s.get('year', 'N/A'),
            'Present ✅',
            now.strftime('%H:%M:%S')
        ])
        r = ws.max_row
        for cell in ws[r]:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=7).fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        ws.cell(row=r, column=7).font = Font(bold=True, color="006100")

    widths = [7, 15, 25, 12, 10, 8, 15, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, len(matched_students)